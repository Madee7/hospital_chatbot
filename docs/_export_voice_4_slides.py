"""Build 4 keyword-focused Split Voice slides for the PPT designer.

Outputs:
  docs/pdf/VOICE_4_SLIDES.pdf
  docs/pdf/VOICE_4_SLIDES.pptx
"""
from __future__ import annotations

import subprocess
import time
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt, Emu

docs = Path(__file__).resolve().parent
out_dir = docs / "pdf"
html_dir = docs / "_pdf_tmp"
out_dir.mkdir(exist_ok=True)
html_dir.mkdir(exist_ok=True)

CHROME = Path(r"C:\Program Files\Google\Chrome\Application\chrome.exe")

TEAL = RGBColor(15, 118, 110)
TEAL_DK = RGBColor(19, 78, 74)
INK = RGBColor(30, 58, 95)
MUTED = RGBColor(100, 116, 139)
WHITE = RGBColor(255, 255, 255)
BG = RGBColor(248, 250, 252)
SOFT = RGBColor(204, 251, 241)
ACCENT = RGBColor(14, 165, 233)
ORANGE = RGBColor(234, 88, 12)
PURPLE = RGBColor(124, 58, 237)


def font(run, size=18, bold=False, color=INK):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = "Segoe UI"


def bg_slide(slide):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = BG
    shape.line.fill.background()


def title_block(slide, title, subtitle=""):
    box = slide.shapes.add_textbox(Inches(0.55), Inches(0.35), Inches(12.2), Inches(0.65))
    tf = box.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = title
    font(r, 30, True, TEAL_DK)
    if subtitle:
        box2 = slide.shapes.add_textbox(Inches(0.55), Inches(0.95), Inches(12.2), Inches(0.4))
        tf2 = box2.text_frame
        tf2.clear()
        p2 = tf2.paragraphs[0]
        r2 = p2.add_run()
        r2.text = subtitle
        font(r2, 14, False, MUTED)


def kw_card(slide, left, top, w, h, keyword, hint, color=TEAL):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(w), Inches(h)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = WHITE
    shape.line.color.rgb = color
    shape.line.width = Pt(2)

    # top accent
    bar = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(left), Inches(top), Inches(w), Inches(0.12)
    )
    bar.fill.solid()
    bar.fill.fore_color.rgb = color
    bar.line.fill.background()

    box = slide.shapes.add_textbox(
        Inches(left + 0.15), Inches(top + 0.35), Inches(w - 0.3), Inches(h - 0.5)
    )
    tf = box.text_frame
    tf.word_wrap = True
    tf.clear()
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = keyword
    font(r, 18, True, color)
    if hint:
        p2 = tf.add_paragraph()
        p2.alignment = PP_ALIGN.CENTER
        r2 = p2.add_run()
        r2.text = hint
        font(r2, 11, False, MUTED)
        p2.space_before = Pt(8)


def pill(slide, left, top, text, w=2.0):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(w), Inches(0.4)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = SOFT
    shape.line.color.rgb = TEAL
    tf = shape.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = text
    font(r, 11, True, TEAL_DK)


def flow_step(slide, left, top, w, h, n, word, color=TEAL):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(w), Inches(h)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = WHITE
    shape.line.color.rgb = color
    shape.line.width = Pt(2.5)

    badge = slide.shapes.add_shape(
        MSO_SHAPE.OVAL, Inches(left + w / 2 - 0.22), Inches(top + 0.25), Inches(0.44), Inches(0.44)
    )
    badge.fill.solid()
    badge.fill.fore_color.rgb = color
    badge.line.fill.background()
    btf = badge.text_frame
    btf.clear()
    bp = btf.paragraphs[0]
    bp.alignment = PP_ALIGN.CENTER
    br = bp.add_run()
    br.text = str(n)
    font(br, 14, True, WHITE)

    box = slide.shapes.add_textbox(
        Inches(left + 0.1), Inches(top + 0.9), Inches(w - 0.2), Inches(h - 1.1)
    )
    tf = box.text_frame
    tf.word_wrap = True
    tf.clear()
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = word
    font(r, 16, True, color)


def arrow(slide, left, top):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.RIGHT_ARROW, Inches(left), Inches(top), Inches(0.32), Inches(0.2)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(148, 163, 184)
    shape.line.fill.background()


def build_pptx() -> Path:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    # ---- 1. WHAT (keywords only) ----
    s = prs.slides.add_slide(blank)
    bg_slide(s)
    title_block(s, "Split Voice System", "Keywords")
    items = [
        (0.5, "Voice Note", "WhatsApp audio", TEAL),
        (3.7, "STT", "Whisper", ACCENT),
        (6.9, "AI Plan", "Planner", PURPLE),
        (10.1, "DB Check", "Executor", ORANGE),
    ]
    for left, kw, hint, color in items:
        kw_card(s, left, 2.0, 2.8, 2.4, kw, hint, color)
    pill(s, 0.5, 5.0, "No TTS", 1.6)
    pill(s, 2.3, 5.0, "Same funnel", 2.0)
    pill(s, 4.5, 5.0, "AR + EN", 1.6)

    # ---- 2. WHY (problem / solve keywords) ----
    s = prs.slides.add_slide(blank)
    bg_slide(s)
    title_block(s, "Why Voice?", "Problem → Solution")

    # problem column
    left_box = s.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.5), Inches(1.6), Inches(5.9), Inches(4.4)
    )
    left_box.fill.solid()
    left_box.fill.fore_color.rgb = WHITE
    left_box.line.color.rgb = ORANGE
    left_box.line.width = Pt(2)
    t = s.shapes.add_textbox(Inches(0.8), Inches(1.85), Inches(5.3), Inches(0.5))
    tf = t.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = "PROBLEM"
    font(r, 14, True, ORANGE)
    for i, kw in enumerate(["Typing friction", "Dialect hard to type", "Chat drop-off", "Slow booking"]):
        row = s.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.85), Inches(2.55 + i * 0.8), Inches(5.2), Inches(0.65)
        )
        row.fill.solid()
        row.fill.fore_color.rgb = RGBColor(255, 247, 237)
        row.line.color.rgb = ORANGE
        tf = row.text_frame
        tf.clear()
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = kw
        font(r, 16, True, ORANGE)

    right_box = s.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.9), Inches(1.6), Inches(5.9), Inches(4.4)
    )
    right_box.fill.solid()
    right_box.fill.fore_color.rgb = WHITE
    right_box.line.color.rgb = TEAL
    right_box.line.width = Pt(2)
    t = s.shapes.add_textbox(Inches(7.2), Inches(1.85), Inches(5.3), Inches(0.5))
    tf = t.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = "SOLUTION"
    font(r, 14, True, TEAL)
    for i, kw in enumerate(["Speak naturally", "Same booking flow", "Faster engage", "No separate app"]):
        row = s.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7.25), Inches(2.55 + i * 0.8), Inches(5.2), Inches(0.65)
        )
        row.fill.solid()
        row.fill.fore_color.rgb = SOFT
        row.line.color.rgb = TEAL
        tf = row.text_frame
        tf.clear()
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = kw
        font(r, 16, True, TEAL)

    # ---- 3. ARCHITECTURE diagram keywords ----
    s = prs.slides.add_slide(blank)
    bg_slide(s)
    title_block(s, "Architecture", "Whisper → Planner → Executor → Reply")
    steps = [
        (0.4, 1, "Whisper", TEAL),
        (3.5, 2, "Planner", ACCENT),
        (6.6, 3, "Executor", TEAL_DK),
        (9.7, 4, "Reply", ORANGE),
    ]
    for left, n, word, color in steps:
        flow_step(s, left, 2.0, 2.8, 2.6, n, word, color)
    for left in (3.2, 6.3, 9.4):
        arrow(s, left, 3.15)
    # keyword row under
    under = [
        (0.4, "STT"),
        (3.5, "VoicePlan"),
        (6.6, "DB Validate"),
        (9.7, "Buttons"),
    ]
    for left, text in under:
        pill(s, left + 0.4, 5.0, text, 2.0)
    note = s.shapes.add_textbox(Inches(0.5), Inches(5.7), Inches(12), Inches(0.5))
    tf = note.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = "AI suggests   ·   System confirms   ·   DB = source of truth"
    font(r, 14, True, TEAL_DK)

    # ---- 4. KEY POINTS ----
    s = prs.slides.add_slide(blank)
    bg_slide(s)
    title_block(s, "Key Points", "Remember these")
    points = [
        (0.45, "Input only", "Not a new bot", TEAL),
        (3.65, "AI plans", "Next booking step", ACCENT),
        (6.85, "System proves", "Live hospital DB", PURPLE),
        (10.05, "Buttons out", "No robot voice", ORANGE),
    ]
    for left, kw, hint, color in points:
        kw_card(s, left, 2.0, 2.85, 2.6, kw, hint, color)
    pill(s, 3.5, 5.2, "Safe", 1.5)
    pill(s, 5.2, 5.2, "Arabic + English", 2.4)
    pill(s, 7.8, 5.2, "One product", 2.0)

    path = out_dir / "VOICE_4_SLIDES.pptx"
    prs.save(path)
    print("OK", path.name, path.stat().st_size)
    return path


PDF_CSS = """
@page { size: A4 landscape; margin: 0; }
* { box-sizing: border-box; }
body { margin: 0; font-family: "Segoe UI", Tahoma, Arial, sans-serif; color: #1e3a5f; }
.slide {
  width: 297mm; height: 210mm; padding: 12mm 14mm;
  page-break-after: always; background: #f8fafc;
}
.slide:last-child { page-break-after: auto; }
h1 { margin: 0; font-size: 28pt; color: #134e4a; }
.sub { margin: 2mm 0 8mm; font-size: 13pt; color: #64748b; }
.grid4 { display: grid; grid-template-columns: 1fr 1fr 1fr 1fr; gap: 5mm; }
.grid2 { display: grid; grid-template-columns: 1fr 1fr; gap: 8mm; }
.card {
  background: #fff; border: 2px solid #0f766e; border-radius: 12px;
  padding: 8mm 4mm; text-align: center; min-height: 42mm;
  border-top: 5px solid #0f766e;
}
.card h2 { margin: 0 0 3mm; font-size: 18pt; color: #0f766e; }
.card p { margin: 0; font-size: 11pt; color: #64748b; }
.card.c2 { border-color: #0ea5e9; border-top-color: #0ea5e9; }
.card.c2 h2 { color: #0ea5e9; }
.card.c3 { border-color: #7c3aed; border-top-color: #7c3aed; }
.card.c3 h2 { color: #7c3aed; }
.card.c4 { border-color: #ea580c; border-top-color: #ea580c; }
.card.c4 h2 { color: #ea580c; }
.kwrow {
  background: #fff; border: 2px solid #ea580c; border-radius: 10px;
  padding: 3.5mm; text-align: center; margin: 3mm 0;
  font-size: 15pt; font-weight: 700; color: #ea580c;
}
.kwrow.ok { border-color: #0f766e; color: #0f766e; }
.col {
  background: #fff; border: 2px solid #cbd5e1; border-radius: 12px; padding: 5mm;
}
.col h3 { margin: 0 0 4mm; font-size: 13pt; letter-spacing: .08em; }
.col.prob { border-color: #ea580c; }
.col.prob h3 { color: #ea580c; }
.col.sol { border-color: #0f766e; }
.col.sol h3 { color: #0f766e; }
.pills { display: flex; gap: 3mm; flex-wrap: wrap; margin-top: 8mm; justify-content: center; }
.pill {
  background: #ccfbf1; border: 1px solid #99f6e4; color: #134e4a;
  padding: 2mm 5mm; border-radius: 999px; font-weight: 700; font-size: 11pt;
}
.flow { display: flex; align-items: center; gap: 3mm; margin-top: 6mm; }
.step {
  flex: 1; background: #fff; border: 2.5px solid #0f766e; border-radius: 14px;
  padding: 8mm 3mm; text-align: center; min-height: 55mm;
}
.step.s2 { border-color: #0ea5e9; }
.step.s3 { border-color: #134e4a; }
.step.s4 { border-color: #ea580c; }
.num {
  width: 34px; height: 34px; border-radius: 50%; background: #0f766e; color: #fff;
  display: inline-flex; align-items: center; justify-content: center;
  font-weight: 700; font-size: 14pt; margin-bottom: 4mm;
}
.step.s2 .num { background: #0ea5e9; }
.step.s3 .num { background: #134e4a; }
.step.s4 .num { background: #ea580c; }
.step h2 { margin: 0; font-size: 18pt; color: #0f766e; }
.step.s2 h2 { color: #0ea5e9; }
.step.s3 h2 { color: #134e4a; }
.step.s4 h2 { color: #ea580c; }
.arrow { font-size: 24pt; color: #94a3b8; font-weight: 700; }
.caption {
  text-align: center; margin-top: 8mm; font-size: 14pt; font-weight: 700; color: #134e4a;
}
"""


def build_pdf() -> Path:
    html = f"""<!DOCTYPE html><html><head><meta charset="utf-8"><style>{PDF_CSS}</style></head><body>

<section class="slide">
  <h1>Split Voice System</h1>
  <p class="sub">Keywords</p>
  <div class="grid4">
    <div class="card"><h2>Voice Note</h2><p>WhatsApp audio</p></div>
    <div class="card c2"><h2>STT</h2><p>Whisper</p></div>
    <div class="card c3"><h2>AI Plan</h2><p>Planner</p></div>
    <div class="card c4"><h2>DB Check</h2><p>Executor</p></div>
  </div>
  <div class="pills">
    <span class="pill">No TTS</span>
    <span class="pill">Same funnel</span>
    <span class="pill">AR + EN</span>
  </div>
</section>

<section class="slide">
  <h1>Why Voice?</h1>
  <p class="sub">Problem → Solution</p>
  <div class="grid2">
    <div class="col prob">
      <h3>PROBLEM</h3>
      <div class="kwrow">Typing friction</div>
      <div class="kwrow">Dialect hard to type</div>
      <div class="kwrow">Chat drop-off</div>
      <div class="kwrow">Slow booking</div>
    </div>
    <div class="col sol">
      <h3>SOLUTION</h3>
      <div class="kwrow ok">Speak naturally</div>
      <div class="kwrow ok">Same booking flow</div>
      <div class="kwrow ok">Faster engage</div>
      <div class="kwrow ok">No separate app</div>
    </div>
  </div>
</section>

<section class="slide">
  <h1>Architecture</h1>
  <p class="sub">Whisper → Planner → Executor → Reply</p>
  <div class="flow">
    <div class="step"><div class="num">1</div><h2>Whisper</h2></div>
    <div class="arrow">→</div>
    <div class="step s2"><div class="num">2</div><h2>Planner</h2></div>
    <div class="arrow">→</div>
    <div class="step s3"><div class="num">3</div><h2>Executor</h2></div>
    <div class="arrow">→</div>
    <div class="step s4"><div class="num">4</div><h2>Reply</h2></div>
  </div>
  <div class="pills">
    <span class="pill">STT</span>
    <span class="pill">VoicePlan</span>
    <span class="pill">DB Validate</span>
    <span class="pill">Buttons</span>
  </div>
  <div class="caption">AI suggests · System confirms · DB = source of truth</div>
</section>

<section class="slide">
  <h1>Key Points</h1>
  <p class="sub">Remember these</p>
  <div class="grid4">
    <div class="card"><h2>Input only</h2><p>Not a new bot</p></div>
    <div class="card c2"><h2>AI plans</h2><p>Next booking step</p></div>
    <div class="card c3"><h2>System proves</h2><p>Live hospital DB</p></div>
    <div class="card c4"><h2>Buttons out</h2><p>No robot voice</p></div>
  </div>
  <div class="pills">
    <span class="pill">Safe</span>
    <span class="pill">Arabic + English</span>
    <span class="pill">One product</span>
  </div>
</section>

</body></html>"""
    html_path = html_dir / "VOICE_4_SLIDES.html"
    html_path.write_text(html, encoding="utf-8")
    pdf_path = out_dir / "VOICE_4_SLIDES.pdf"
    if pdf_path.exists():
        try:
            pdf_path.unlink()
        except OSError:
            pass
    subprocess.run(
        [
            str(CHROME),
            "--headless=new",
            "--disable-gpu",
            "--no-pdf-header-footer",
            f"--print-to-pdf={pdf_path}",
            html_path.resolve().as_uri(),
        ],
        check=False,
    )
    for _ in range(40):
        if pdf_path.exists() and pdf_path.stat().st_size > 500:
            break
        time.sleep(0.25)
    print("OK", pdf_path.name, pdf_path.stat().st_size if pdf_path.exists() else 0)
    return pdf_path


def main() -> None:
    build_pptx()
    if CHROME.exists():
        build_pdf()


if __name__ == "__main__":
    main()
