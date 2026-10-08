"""$1 table in the same style as the existing PPT cost slide.

Outputs:
  docs/pdf/COST_ONE_DOLLAR.pdf
  docs/pdf/COST_ONE_DOLLAR.pptx
  docs/images/cost-one-dollar.png
"""
from __future__ import annotations

import subprocess
import time
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt

docs = Path(__file__).resolve().parent
out_dir = docs / "pdf"
img_dir = docs / "images"
html_dir = docs / "_pdf_tmp"
out_dir.mkdir(exist_ok=True)
img_dir.mkdir(exist_ok=True)
html_dir.mkdir(exist_ok=True)

CHROME = Path(r"C:\Program Files\Google\Chrome\Application\chrome.exe")

# Sampled from the existing PPT table screenshot
SLIDE_BG = (233, 240, 250)
HEADER = (74, 86, 112)
ROW_ODD = (233, 237, 246)
ROW_EVEN = (222, 228, 240)
INK = (45, 55, 72)
WHITE = (255, 255, 255)
GRID = (210, 218, 230)

HEADER_RGB = RGBColor(*HEADER)
INK_RGB = RGBColor(*INK)
BG_RGB = RGBColor(*SLIDE_BG)

# Same 5-row shape as the slide: Metric | Cost | Basis
ROWS = [
    ("Complete voice notes", "≈ 1,170", "Whisper + GPT-4o-mini (115 voice notes)"),
    ("LLM voice planner only", "≈ 2,160", "gpt-4o-mini voice_parse"),
    ("Whisper STT only", "≈ 2,560", "whisper-1"),
    ("Expected voice compute", "≈ 6–18 s", "Typical end-to-end reply (≈ 4 s of audio)"),
    ("Test volume", "1,004 WhatsApp messages", "790 inbound · 214 outbound · 0 failed"),
]
HEADERS = ["Metric", "Cost", "Basis"]


def font(size: int, bold: bool = False):
    candidates = [
        ("C:/Windows/Fonts/calibrib.ttf" if bold else "C:/Windows/Fonts/calibri.ttf"),
        ("C:/Windows/Fonts/segoeuib.ttf" if bold else "C:/Windows/Fonts/segoeui.ttf"),
        ("C:/Windows/Fonts/arialbd.ttf" if bold else "C:/Windows/Fonts/arial.ttf"),
    ]
    for path in candidates:
        try:
            return ImageFont.truetype(path, size)
        except OSError:
            continue
    return ImageFont.load_default()


def render_png() -> Path:
    """Wide short table, same proportions as the PPT screenshot (1024×382)."""
    w, h = 1920, 716
    img = Image.new("RGB", (w, h), SLIDE_BG)
    draw = ImageDraw.Draw(img)
    f_h = font(28, True)
    f_body = font(26, False)

    margin = 48
    x0, y0 = margin, 56
    x1, y1 = w - margin, h - 56
    header_h = 78
    row_h = (y1 - y0 - header_h) / len(ROWS)
    # column splits match the original: metric ~32%, cost ~28%, basis rest
    c0, c1, c2, c3 = x0, x0 + int((x1 - x0) * 0.32), x0 + int((x1 - x0) * 0.60), x1
    cols = [c0, c1, c2]

    draw.rounded_rectangle([x0, y0, x1, y1], radius=10, fill=ROW_ODD)

    # header
    draw.rectangle([x0, y0, x1, y0 + header_h], fill=HEADER)
    # clip header to rounded top by redrawing corners in header color
    draw.pieslice([x0, y0, x0 + 20, y0 + 20], 180, 270, fill=HEADER)
    draw.pieslice([x1 - 20, y0, x1, y0 + 20], 270, 360, fill=HEADER)

    for i, label in enumerate(HEADERS):
        draw.text((cols[i] + 28, y0 + 22), label, font=f_h, fill=WHITE)

    y = y0 + header_h
    for i, row in enumerate(ROWS):
        fill = ROW_ODD if i % 2 == 0 else ROW_EVEN
        bottom = y1 if i == len(ROWS) - 1 else y + row_h
        draw.rectangle([x0, y, x1, bottom], fill=fill)
        draw.line([x0, y, x1, y], fill=GRID, width=1)
        for j, text in enumerate(row):
            draw.text((cols[j] + 28, y + 22), text, font=f_body, fill=INK)
        y = bottom

    # outer stroke
    draw.rounded_rectangle([x0, y0, x1, y1], radius=10, outline=HEADER, width=2)

    path = img_dir / "cost-one-dollar.png"
    img.save(path, "PNG")
    print("OK", path, path.stat().st_size)
    return path


def export_pdf(png: Path) -> None:
    html = f"""<!DOCTYPE html>
<html><head><meta charset="utf-8">
<style>
@page {{ size: A4 landscape; margin: 0; }}
html, body {{
  margin: 0;
  padding: 0;
  background: rgb(233, 240, 250);
}}
img {{
  width: 100%;
  height: 100vh;
  object-fit: contain;
  display: block;
}}
</style></head>
<body>
<img src="{png.resolve().as_uri()}">
</body></html>
"""
    html_path = html_dir / "COST_ONE_DOLLAR.html"
    html_path.write_text(html, encoding="utf-8")
    pdf_path = out_dir / "COST_ONE_DOLLAR.pdf"
    if pdf_path.exists():
        try:
            pdf_path.unlink()
        except OSError:
            pass
    subprocess.run(
        [
            str(CHROME),
            "--headless=new",
            "--disable-gpu",
            "--no-pdf-header-footer",
            f"--print-to-pdf={pdf_path}",
            html_path.resolve().as_uri(),
        ],
        check=False,
    )
    for _ in range(30):
        if pdf_path.exists() and pdf_path.stat().st_size > 1000:
            break
        time.sleep(0.2)
    print("OK", pdf_path, pdf_path.stat().st_size if pdf_path.exists() else 0)


def export_pptx() -> None:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    bg = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5)
    )
    bg.fill.solid()
    bg.fill.fore_color.rgb = BG_RGB
    bg.line.fill.background()

    table_shape = slide.shapes.add_table(
        1 + len(ROWS), 3, Inches(0.7), Inches(1.55), Inches(11.9), Inches(4.4)
    )
    table = table_shape.table
    table.columns[0].width = Inches(3.8)
    table.columns[1].width = Inches(3.3)
    table.columns[2].width = Inches(4.8)

    data = [HEADERS] + [list(row) for row in ROWS]
    for i, row in enumerate(data):
        for j, text in enumerate(row):
            cell = table.cell(i, j)
            cell.text = ""
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            p = cell.text_frame.paragraphs[0]
            p.alignment = PP_ALIGN.LEFT
            run = p.add_run()
            run.text = text
            run.font.name = "Calibri"
            run.font.size = Pt(18 if i else 20)
            run.font.bold = i == 0
            cell.fill.solid()
            if i == 0:
                run.font.color.rgb = RGBColor(255, 255, 255)
                cell.fill.fore_color.rgb = HEADER_RGB
            else:
                run.font.color.rgb = INK_RGB
                shade = ROW_ODD if i % 2 == 1 else ROW_EVEN
                cell.fill.fore_color.rgb = RGBColor(*shade)

    path = out_dir / "COST_ONE_DOLLAR.pptx"
    prs.save(path)
    print("OK", path, path.stat().st_size)


def main() -> None:
    png = render_png()
    export_pdf(png)
    export_pptx()


if __name__ == "__main__":
    main()
