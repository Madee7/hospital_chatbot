"""Export Hospital Chatbot presentation to PowerPoint.

Output:
  docs/pdf/HOSPITAL_CHATBOT.pptx
"""
from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt

docs = Path(__file__).resolve().parent
out_dir = docs / "pdf"
out_dir.mkdir(exist_ok=True)

# Theme (matches Yor deck: slate / soft blue)
SLATE = RGBColor(77, 89, 115)
SLATE_DK = RGBColor(48, 53, 64)
INK = RGBColor(48, 53, 64)
MUTED = RGBColor(100, 116, 139)
WHITE = RGBColor(255, 255, 255)
BG = RGBColor(245, 247, 251)
SOFT = RGBColor(214, 222, 241)
ACCENT = RGBColor(77, 89, 115)
BORDER = RGBColor(188, 196, 215)
TEAL = RGBColor(15, 118, 110)
ORANGE = RGBColor(234, 88, 12)
PURPLE = RGBColor(99, 102, 241)


def font(run, size=18, bold=False, color=INK):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = "Segoe UI"


def bg_slide(slide, color=BG):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()


def section_slide(prs, blank, number: str, title: str):
    s = prs.slides.add_slide(blank)
    bg_slide(s, WHITE)
    box = s.shapes.add_textbox(Inches(0.8), Inches(2.4), Inches(11.5), Inches(1.0))
    tf = box.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = number
    font(r, 54, True, SLATE)
    box2 = s.shapes.add_textbox(Inches(0.8), Inches(3.5), Inches(11.5), Inches(1.2))
    tf2 = box2.text_frame
    tf2.word_wrap = True
    tf2.clear()
    p2 = tf2.paragraphs[0]
    r2 = p2.add_run()
    r2.text = title
    font(r2, 36, True, SLATE_DK)
    # accent line
    line = s.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(5.0), Inches(3.5), Inches(0.06)
    )
    line.fill.solid()
    line.fill.fore_color.rgb = SLATE
    line.line.fill.background()
    return s


def title_block(slide, title, subtitle=""):
    box = slide.shapes.add_textbox(Inches(0.55), Inches(0.3), Inches(12.2), Inches(0.6))
    tf = box.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = title
    font(r, 28, True, SLATE_DK)
    if subtitle:
        box2 = slide.shapes.add_textbox(Inches(0.55), Inches(0.9), Inches(12.2), Inches(0.45))
        tf2 = box2.text_frame
        tf2.word_wrap = True
        tf2.clear()
        p2 = tf2.paragraphs[0]
        r2 = p2.add_run()
        r2.text = subtitle
        font(r2, 14, False, MUTED)


def card(slide, left, top, w, h, title, bullets, accent=SLATE):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(w), Inches(h)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = WHITE
    shape.line.color.rgb = BORDER
    shape.line.width = Pt(1.5)

    bar = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(left), Inches(top), Inches(w), Inches(0.1)
    )
    bar.fill.solid()
    bar.fill.fore_color.rgb = accent
    bar.line.fill.background()

    box = slide.shapes.add_textbox(
        Inches(left + 0.2), Inches(top + 0.25), Inches(w - 0.4), Inches(h - 0.4)
    )
    tf = box.text_frame
    tf.word_wrap = True
    tf.clear()
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = title
    font(r, 16, True, accent)
    for b in bullets:
        p2 = tf.add_paragraph()
        r2 = p2.add_run()
        r2.text = f"• {b}"
        font(r2, 13, False, INK)
        p2.space_before = Pt(6)


def step_row(slide, left, top, items, w=2.8, h=1.6, gap=0.25):
    for i, (title, desc) in enumerate(items):
        x = left + i * (w + gap)
        shape = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(top), Inches(w), Inches(h)
        )
        shape.fill.solid()
        shape.fill.fore_color.rgb = SOFT
        shape.line.color.rgb = BORDER

        badge = slide.shapes.add_shape(
            MSO_SHAPE.OVAL, Inches(x + 0.15), Inches(top + 0.2), Inches(0.38), Inches(0.38)
        )
        badge.fill.solid()
        badge.fill.fore_color.rgb = SLATE
        badge.line.fill.background()
        btf = badge.text_frame
        btf.clear()
        bp = btf.paragraphs[0]
        bp.alignment = PP_ALIGN.CENTER
        br = bp.add_run()
        br.text = str(i + 1)
        font(br, 12, True, WHITE)

        box = slide.shapes.add_textbox(
            Inches(x + 0.15), Inches(top + 0.7), Inches(w - 0.3), Inches(h - 0.85)
        )
        tf = box.text_frame
        tf.word_wrap = True
        tf.clear()
        p = tf.paragraphs[0]
        r = p.add_run()
        r.text = title
        font(r, 14, True, SLATE_DK)
        if desc:
            p2 = tf.add_paragraph()
            r2 = p2.add_run()
            r2.text = desc
            font(r2, 11, False, MUTED)
            p2.space_before = Pt(4)


def bullet_slide(slide, items, left=0.7, top=1.5, w=12.0):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(w), Inches(5.5))
    tf = box.text_frame
    tf.word_wrap = True
    tf.clear()
    first = True
    for title, desc in items:
        if first:
            p = tf.paragraphs[0]
            first = False
        else:
            p = tf.add_paragraph()
            p.space_before = Pt(14)
        r = p.add_run()
        r.text = title
        font(r, 18, True, SLATE_DK)
        if desc:
            p2 = tf.add_paragraph()
            r2 = p2.add_run()
            r2.text = desc
            font(r2, 13, False, MUTED)
            p2.space_before = Pt(4)


def add_table(slide, rows, left=0.55, top=1.5, width=12.2, row_h=0.55):
    cols = len(rows[0])
    table_shape = slide.shapes.add_table(len(rows), cols, Inches(left), Inches(top), Inches(width), Inches(row_h * len(rows)))
    table = table_shape.table
    col_w = int(Inches(width) / cols)
    for i in range(cols):
        table.columns[i].width = col_w
    for r_idx, row in enumerate(rows):
        for c_idx, cell_text in enumerate(row):
            cell = table.cell(r_idx, c_idx)
            cell.text = cell_text
            for p in cell.text_frame.paragraphs:
                for run in p.runs:
                    font(run, 12, bold=(r_idx == 0), color=WHITE if r_idx == 0 else INK)
            # header / body fill via cell fill is limited; skip fancy fills


def build_pptx() -> Path:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    # ---- Title ----
    s = prs.slides.add_slide(blank)
    bg_slide(s, WHITE)
    box = s.shapes.add_textbox(Inches(0.8), Inches(2.5), Inches(11.7), Inches(1.0))
    tf = box.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = "Hospital Chatbot"
    font(r, 44, True, SLATE_DK)
    line = s.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(5.4), Inches(3.7), Inches(2.5), Inches(0.05)
    )
    line.fill.solid()
    line.fill.fore_color.rgb = SLATE
    line.line.fill.background()
    box2 = s.shapes.add_textbox(Inches(0.8), Inches(4.0), Inches(11.7), Inches(1.0))
    tf2 = box2.text_frame
    tf2.word_wrap = True
    tf2.clear()
    p2 = tf2.paragraphs[0]
    p2.alignment = PP_ALIGN.CENTER
    r2 = p2.add_run()
    r2.text = "By: Eng / Ghena Ahmed  ·  Eng / Mohamed Madeeh  ·  Eng / Yousof Ahmed"
    font(r2, 14, False, MUTED)

    # ---- 01 Why WhatsApp ----
    section_slide(prs, blank, "01", "Why WhatsApp for Healthcare?")

    s = prs.slides.add_slide(blank)
    bg_slide(s)
    title_block(s, "Why WhatsApp for Healthcare")
    card(
        s, 0.7, 1.6, 5.8, 4.5, "Patients",
        ["24/7 Access", "Faster Booking", "Easy Navigation"],
        TEAL,
    )
    card(
        s, 6.9, 1.6, 5.8, 4.5, "Reception",
        ["Less Manual Work", "Fewer Calls", "Fewer Errors"],
        ORANGE,
    )

    # ---- 02 Solution Architecture ----
    section_slide(prs, blank, "02", "Solution Architecture")

    s = prs.slides.add_slide(blank)
    bg_slide(s)
    title_block(
        s,
        "Main Menu: Single Entry Point",
        "One centralized interface handles all patient requests with clear navigation.",
    )
    items = [
        ("Book Appointment", "Start guided scheduling with real availability"),
        ("Reschedule", "Change date or time instantly"),
        ("Cancel", "Release the slot immediately"),
        ("Ask about Doctors/Specialties", "Browse provider information"),
        ("Talk to Reception", "Human handover for complex cases"),
    ]
    bullet_slide(s, items, top=1.6)

    # ---- 03 Booking Journey ----
    section_slide(prs, blank, "03", "Booking Journey")

    s = prs.slides.add_slide(blank)
    bg_slide(s)
    title_block(s, "Booking Journey", "Select Appointment")
    step_row(
        s,
        0.8,
        2.2,
        [
            ("Start Booking", ""),
            ("Select Doctor", ""),
            ("Select Date & Time", ""),
        ],
        w=3.6,
        h=2.2,
        gap=0.4,
    )

    s = prs.slides.add_slide(blank)
    bg_slide(s)
    title_block(s, "Booking Journey", "Review & Confirm")
    step_row(
        s,
        0.8,
        2.2,
        [
            ("Enter Full Name", ""),
            ("Review Booking", ""),
            ("Confirm Appointment", ""),
        ],
        w=3.6,
        h=2.2,
        gap=0.4,
    )

    s = prs.slides.add_slide(blank)
    bg_slide(s)
    title_block(
        s,
        "Flexible Search Paths",
        "Two distinct pathways converge on the same outcome, accommodating patient preferences.",
    )
    card(s, 0.7, 1.8, 5.8, 3.8, "By Specialty", ["Select specialty first, then choose doctor"], PURPLE)
    card(s, 6.9, 1.8, 5.8, 3.8, "By Doctor Name", ["Browse doctors directly"], TEAL)

    s = prs.slides.add_slide(blank)
    bg_slide(s)
    title_block(
        s,
        "Confirmation with Conflict Protection",
        "Multi-layer safeguards hold slots and prevent simultaneous double-bookings.",
    )
    step_row(
        s,
        0.45,
        1.8,
        [
            ("Temporary Slot Hold", "Slot reserved during confirmation window"),
            ("Review Booking Details", "Doctor, date, time & personal info"),
            ("Final Availability Check", "Verify slot before saving"),
            ("Result", "Secure confirmation, no overbooking"),
        ],
        w=3.0,
        h=3.8,
        gap=0.2,
    )

    # ---- 04 Reschedule ----
    section_slide(prs, blank, "04", "Reschedule Appointment")

    s = prs.slides.add_slide(blank)
    bg_slide(s)
    title_block(s, "Reschedule Appointment")
    step_row(
        s,
        0.8,
        1.8,
        [
            ("Select Appointment", "Patient selects an active appointment to modify."),
            ("Choose Update Type", "Change Date, Time, or Doctor"),
            ("Dynamic Update Flow", "Date → dates · Time → slots · Doctor → list"),
        ],
        w=3.6,
        h=3.5,
        gap=0.4,
    )

    s = prs.slides.add_slide(blank)
    bg_slide(s)
    title_block(s, "Reschedule Appointment", "Complete the update")
    step_row(
        s,
        0.8,
        1.8,
        [
            ("Select New Slot", "Pick from live available schedules."),
            ("Review & Confirm", 'Patient reviews summary and clicks "Confirm".'),
            ("Live Sync & Receipt", "DB updates instantly + Maps link."),
        ],
        w=3.6,
        h=3.5,
        gap=0.4,
    )

    # ---- 05 Cancel ----
    section_slide(prs, blank, "05", "Cancel Appointment")

    s = prs.slides.add_slide(blank)
    bg_slide(s)
    title_block(s, "Instant Cancellation")
    step_row(
        s,
        0.45,
        1.8,
        [
            ("Menu Selection", 'Tap "Cancel appointment" or enter Ref ID.'),
            ("Booking Selection", "Choose from active upcoming appointments."),
            ("Confirmation", "Confirm cancellation in one step."),
            ("Slot Release", "Reserved slot freed immediately."),
        ],
        w=3.0,
        h=3.8,
        gap=0.2,
    )

    # ---- 06 Reception Handover ----
    section_slide(prs, blank, "06", "Reception Handover")

    s = prs.slides.add_slide(blank)
    bg_slide(s)
    title_block(
        s,
        "Reception Handover Protocol",
        "Complex cases smoothly escalate to human staff with full context preservation.",
    )
    cards = [
        (0.5, "Main Menu Option", '"Talk to reception" available at every decision point'),
        (3.7, "Context Preservation", "Full conversation history transferred to staff"),
        (6.9, "Notification", "Reception can be alerted when handover is triggered"),
        (10.1, "Safety Net", "No patient inquiry goes unanswered"),
    ]
    for left, title, desc in cards:
        card(s, left, 1.8, 2.9, 4.0, title, [desc], SLATE)

    # ---- 07 Advanced / Voice ----
    section_slide(prs, blank, "07", "Advanced System Capabilities")

    s = prs.slides.add_slide(blank)
    bg_slide(s)
    title_block(s, "Voice System")
    step_row(
        s,
        0.45,
        1.6,
        [
            ("Voice Note", "Patient audio input starts the flow."),
            ("STT (Whisper)", "Speech transcribed into structured text."),
            ("AI Plan", "Intent recognition maps text to action."),
            ("DB Check", "Live hospital DB verifies availability."),
        ],
        w=3.0,
        h=3.2,
        gap=0.2,
    )
    # pills
    for i, label in enumerate(["No TTS", "Same Funnel", "AR + EN"]):
        shape = s.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            Inches(0.7 + i * 2.4),
            Inches(5.4),
            Inches(2.1),
            Inches(0.55),
        )
        shape.fill.solid()
        shape.fill.fore_color.rgb = SOFT
        shape.line.color.rgb = BORDER
        tf = shape.text_frame
        tf.clear()
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = label
        font(r, 13, True, SLATE_DK)

    s = prs.slides.add_slide(blank)
    bg_slide(s)
    title_block(s, "Why Voice?")
    card(
        s, 0.7, 1.6, 5.5, 4.5, "The problem",
        ["Typing friction", "Dialect hard to type", "Chat drop-off", "Slow booking"],
        ORANGE,
    )
    card(
        s, 7.0, 1.6, 5.5, 4.5, "SOLUTION",
        ["Speak naturally", "Same booking flow", "Faster engage", "No separate app"],
        TEAL,
    )

    s = prs.slides.add_slide(blank)
    bg_slide(s)
    title_block(s, "Architecture")
    step_row(
        s,
        0.45,
        1.6,
        [
            ("Whisper", "STT (Speech to Text)"),
            ("Planner", "VoicePlan / AI Intent"),
            ("Executor", "DB Verification"),
            ("Reply", "WhatsApp Buttons"),
        ],
        w=3.0,
        h=2.8,
        gap=0.2,
    )
    box = s.shapes.add_textbox(Inches(0.7), Inches(5.0), Inches(12), Inches(1.2))
    tf = box.text_frame
    tf.word_wrap = True
    tf.clear()
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = "AI suggests  →  System confirms  →  DB = source of truth"
    font(r, 18, True, SLATE_DK)

    s = prs.slides.add_slide(blank)
    bg_slide(s)
    title_block(s, "AI Cost Results")
    rows = [
        ("Metric", "Cost", "Basis"),
        ("Total AI cost (30 days)", "$0.0931", "Live dashboard logs"),
        ("Per voice message", "≈ $0.0009", "Whisper + GPT-4o-mini (94 voice notes)"),
        ("LLM voice planner only", "≈ $0.0004", "gpt-4o-mini voice_parse"),
        ("Whisper STT only", "≈ $0.0004", "whisper-1"),
        ("Test volume", "783 WhatsApp messages", "617 inbound · 166 outbound · 0 failed"),
    ]
    # Manual table as text cards for cleaner look
    headers = rows[0]
    y = 1.45
    # header bar
    hdr = s.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0.55), Inches(y), Inches(12.2), Inches(0.5)
    )
    hdr.fill.solid()
    hdr.fill.fore_color.rgb = SLATE
    hdr.line.fill.background()
    widths = [3.8, 2.6, 5.8]
    x = 0.55
    for i, h in enumerate(headers):
        box = s.shapes.add_textbox(Inches(x + 0.15), Inches(y + 0.08), Inches(widths[i] - 0.2), Inches(0.4))
        tf = box.text_frame
        tf.clear()
        p = tf.paragraphs[0]
        r = p.add_run()
        r.text = h
        font(r, 13, True, WHITE)
        x += widths[i]
    y += 0.5
    for idx, row in enumerate(rows[1:]):
        bgc = WHITE if idx % 2 == 0 else SOFT
        bar = s.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches(0.55), Inches(y), Inches(12.2), Inches(0.7)
        )
        bar.fill.solid()
        bar.fill.fore_color.rgb = bgc
        bar.line.color.rgb = BORDER
        x = 0.55
        for i, cell in enumerate(row):
            box = s.shapes.add_textbox(Inches(x + 0.15), Inches(y + 0.15), Inches(widths[i] - 0.2), Inches(0.5))
            tf = box.text_frame
            tf.word_wrap = True
            tf.clear()
            p = tf.paragraphs[0]
            r = p.add_run()
            r.text = cell
            font(r, 12, bold=(i == 1), color=INK)
            x += widths[i]
        y += 0.7

    s = prs.slides.add_slide(blank)
    bg_slide(s)
    title_block(s, "Key Points")
    cards = [
        (0.5, "Input only", "not a new bot"),
        (3.7, "AI plans", "next booking step"),
        (6.9, "System proves", "live hospital DB"),
        (10.1, "Buttons out", "no robot voice"),
    ]
    for left, title, desc in cards:
        card(s, left, 1.6, 2.9, 3.2, title, [desc], SLATE)
    box = s.shapes.add_textbox(Inches(0.7), Inches(5.3), Inches(12), Inches(1.0))
    tf = box.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = "Safe        ·        Arabic + English        ·        One product"
    font(r, 18, True, SLATE_DK)

    # ---- 08 Dashboard ----
    section_slide(prs, blank, "08", "Clinic Management Dashboard")

    s = prs.slides.add_slide(blank)
    bg_slide(s)
    title_block(
        s,
        "Clinic Dashboard Capabilities",
        "End-to-end control over clinic setup, medical schedules, and patient operations",
    )
    bullet_slide(
        s,
        [
            ("Clinic Setup & KPIs", "Track daily booking metrics, update clinic details, and manage system settings."),
            ("Doctors & Schedules", "Manage departments, doctor profiles, working hours, and schedules."),
            ("Appointments & Chats", "View bookings, cancel/reschedule, follow up on live conversations."),
            ("Reminders & Costs", "Track automated WhatsApp reminders and monitor real-time AI usage costs."),
        ],
        top=1.6,
    )

    s = prs.slides.add_slide(blank)
    bg_slide(s)
    title_block(s, "Dashboard Overview & Setup", "Real-time KPIs, clinic configuration, and medical registry")
    cards = [
        (0.5, "Operational KPIs & Cost Tracking", ["Live booking & chat stats", "Exact AI costs in real time"]),
        (4.7, "Clinic Setup & Holidays", ["Clinic details & WhatsApp number", "Holiday dates and location link"]),
        (8.9, "Doctor & Specialty Registry", ["Departments & doctor profiles", "Bilingual doctor details"]),
    ]
    for left, title, bullets in cards:
        card(s, left, 1.7, 3.9, 4.2, title, bullets, SLATE)

    s = prs.slides.add_slide(blank)
    bg_slide(s)
    title_block(s, "Doctor Schedules & Shift Control", "Doctor shifts, slot durations, and leave management")
    card(
        s, 0.7, 1.7, 5.8, 4.2, "Working Hours & Doctor Leave",
        ["Weekly doctor shifts & slot durations", "Auto-blocks slots during doctor absence"],
        TEAL,
    )
    card(
        s, 6.9, 1.7, 5.8, 4.2, "Doctor Schedule & Slots",
        ["Tracks free and booked times", "Hides confirmed slots automatically"],
        PURPLE,
    )

    s = prs.slides.add_slide(blank)
    bg_slide(s)
    title_block(s, "Appointments Management", "Live booking table, updates, and data export")
    card(
        s, 0.7, 1.7, 5.8, 4.2, "Appointments Table",
        ["View all patient booking records", "Filter list by booking status"],
        SLATE,
    )
    card(
        s, 6.9, 1.7, 5.8, 4.2, "Cancel & Reschedule",
        ["Cancel or reschedule appointments", "Download data with Export CSV"],
        ORANGE,
    )

    s = prs.slides.add_slide(blank)
    bg_slide(s)
    title_block(s, "Operational Control", "Live conversations and AI cost analytics")
    card(
        s, 0.7, 1.7, 5.8, 4.2, "Conversations & Human Handover",
        ["Live patient chat transcripts", "Open cases for human reception"],
        TEAL,
    )
    card(
        s, 6.9, 1.7, 5.8, 4.2, "AI Usage & Cost Analytics",
        ["Tokens and response speed", "AI models cost breakdown"],
        PURPLE,
    )

    # ---- Thank You ----
    s = prs.slides.add_slide(blank)
    bg_slide(s, WHITE)
    box = s.shapes.add_textbox(Inches(0.8), Inches(2.8), Inches(11.7), Inches(1.2))
    tf = box.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = "Thank You"
    font(r, 48, True, SLATE_DK)
    box2 = s.shapes.add_textbox(Inches(0.8), Inches(4.2), Inches(11.7), Inches(0.6))
    tf2 = box2.text_frame
    tf2.clear()
    p2 = tf2.paragraphs[0]
    p2.alignment = PP_ALIGN.CENTER
    r2 = p2.add_run()
    r2.text = "Hospital Chatbot"
    font(r2, 18, False, MUTED)

    path = out_dir / "HOSPITAL_CHATBOT.pptx"
    prs.save(path)
    print(f"Wrote {path}")
    return path


if __name__ == "__main__":
    build_pptx()
